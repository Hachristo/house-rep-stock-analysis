from openpyxl import load_workbook
import json

def Read(xlsx_path):
    representatives = []
    wb = load_workbook(filename = xlsx_path)
    ws = wb.active
    for row in ws.iter_rows(min_row = 2):
        if row[0].value == None:
            break
        rep = {}
        rep['BioguideID'] = row[0].value
        rep['Completed'] = row[1].value

        rep['ClerkOffice'] = []

        term1 = {}
        term1['YearRange'] = [int(row[3].value), int(row[4].value)]
        term1['LastName'] = row[5].value
        term1['District'] = row[6].value
        rep['ClerkOffice'].append(term1)

        if row[8].value != None:
            term2 = {}
            term2['YearRange'] = [int(row[8].value), int(row[9].value)]
            term2['LastName'] = row[10].value
            term2['District'] = row[11].value
            rep['ClerkOffice'].append(term2)

        if row[13].value != None:
            term3 = {}
            term3['YearRange'] = [int(row[13].value), int(row[14].value)]
            term3['LastName'] = row[15].value
            term3['District'] = row[16].value
            rep['ClerkOffice'].append(term3)

        representatives.append(rep)
    return representatives

def Update(representatives, xlsx_path):
    with open(xlsx_path, "w") as f:
        json.dump(representatives, f, indent=4)

reps = Read('./test.xlsx')
Update(reps, './reps.json')