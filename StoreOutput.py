from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles import Alignment
from openpyxl.styles import Color
import json
import os
import sys

# read json file at path and output as dictionary
def readJsonFile(path):
    with open(path, 'r') as f:
        return json.load(f)

# convert json discrepancy entries into excel sheet
def fill_sheet(worksheet, json):
    ticker_font = Font(size = 16, bold = True)
    link_font = Font(color=Color(rgb="000000FF"), underline="single")
    for year in json.keys():
        for stock in json[year].keys():
            # convert list of PTR links into a single string seperated by newlines
            PTRs = json[year][stock]['PTR']
            single_cell_PTRs = ''
            if PTRs != None:
                single_cell_PTRs = '\n'.join(PTRs)
            
            # convert list of transaction dates into a single string seperated by newlines
            TDs = json[year][stock]['Transaction_Dates']
            single_cell_TDs = ''
            if TDs != None:
                single_cell_TDs = '\n'.join(TDs)
            
            # catch errors in parsing amounts
            if json[year][stock]['Amount'] == 'ERROR':
                lower_bound = 0
                upper_bound = 0
            else:
                lower_bound = json[year][stock]['Amount'][0]
                upper_bound = json[year][stock]['Amount'][1]

            worksheet.append([stock,
                    json[year][stock]['Transaction'],
                    lower_bound,
                    upper_bound,
                    json[year][stock]['Present_In'],
                    json[year][stock]['Missing_In'],
                    json[year][stock]['Number_of_Transactions'],
                    single_cell_TDs,
                    '',
                    '',
                    single_cell_PTRs,
                    year])
            
            # convert previous and current FD cells into clickable links
            current_row = worksheet.max_row
            worksheet['I' + str(current_row)].hyperlink = json[year][stock]['Previous_FD']
            worksheet['I' + str(current_row)].value = year[:4]
            worksheet['I' + str(current_row)].font = link_font
            worksheet['J' + str(current_row)].hyperlink = json[year][stock]['Current_FD']
            worksheet['J' + str(current_row)].value = year[-4:]
            worksheet['J' + str(current_row)].font = link_font

    # increase font size for stock ticker abbreviation
    lowest_row = worksheet.max_row
    for row in range(2, lowest_row + 1):
        worksheet['A' + str(row)].font = ticker_font

# apply custom formatting to improve spreadsheet readability
def format(worksheet):
    # manually fit columns
    worksheet.column_dimensions['A'].width = 11.3
    worksheet.column_dimensions['B'].width = 11.2
    worksheet.column_dimensions['C'].width = 20.9
    worksheet.column_dimensions['D'].width = 20.9
    worksheet.column_dimensions['E'].width = 21.8
    worksheet.column_dimensions['F'].width = 21.8
    worksheet.column_dimensions['G'].width = 17.4
    worksheet.column_dimensions['H'].width = 16.2
    worksheet.column_dimensions['I'].width = 11.3
    worksheet.column_dimensions['J'].width = 11.3
    worksheet.column_dimensions['K'].width = 73.2
    worksheet.column_dimensions['L'].width = 11.3

    # set all cells to be center aligned
    centered_alignment = Alignment(horizontal='center', vertical='center')
    for row in worksheet:
        for cell in row:
            cell.alignment = centered_alignment

def main(name):
    wb = Workbook()
    for entry in os.listdir('./Results'):
        if entry != '.gitignore' and 'trading' not in entry:
            fullpath = os.path.join('./Results', entry)
            try:
                js = readJsonFile(fullpath)
            except json.JSONDecodeError:
                print(f"Output Error: \'" + fullpath + "\' could not be decoded")
                sys.exit(10)
            index = entry.find('_')
            ws = wb.create_sheet(entry[:index])
            ws.append(['Stock',
                       'Transaction',
                       'Amount, Lower Bound',
                       'Amount, Upper Bound',
                       'Present In',
                       'Missing In',
                       'Transaction Count',
                       'Transaction Dates',
                       'Previous FD',
                       'Current FD',
                       'PTRs',
                       'Year'])
            fill_sheet(ws, js)
            format(ws)
    wb.remove(wb['Sheet'])
    wb.save('./Sheets/' + name + '.xlsx')
